"""Tests for optional local NER foundation."""

from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch

from docx import Document


PROJECT_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT_ROOT / "src"))

from anonymizer import anonymize_batch, anonymize_docx_file, anonymize_file
from ner import (
    DEFAULT_NER_MODEL,
    NER_EXCLUSION_LINEBREAK_NON_PERSON,
    NER_STATUS_DEPENDENCY_MISSING,
    NER_STATUS_DISABLED,
    NER_STATUS_MODEL_MISSING,
    detect_entities,
    detect_entities_with_details,
    detect_ner_support,
    prepare_ner_context,
)


def workspace_temp_dir():
    return tempfile.TemporaryDirectory(dir=PROJECT_ROOT / "tests")


class FakeEntity:
    def __init__(self, start: int, end: int, label: str) -> None:
        self.start_char = start
        self.end_char = end
        self.label_ = label


class FakeDoc:
    def __init__(self, ents: list[FakeEntity]) -> None:
        self.ents = ents


class FakeNerModel:
    def __init__(self, rules: list[tuple[str, str]]) -> None:
        self.rules = rules

    def __call__(self, text: str) -> FakeDoc:
        entities: list[FakeEntity] = []
        for value, label in self.rules:
            start = text.find(value)
            if start >= 0:
                entities.append(FakeEntity(start, start + len(value), label))
        return FakeDoc(entities)


class FakeSpacy:
    def __init__(self, model: FakeNerModel | None = None, error: Exception | None = None):
        self.model = model
        self.error = error

    def load(self, model_name: str):
        if self.error is not None:
            raise self.error
        return self.model


class NerFoundationTests(unittest.TestCase):
    def test_ner_detection_reports_disabled_status(self) -> None:
        status = detect_ner_support(enabled=False)

        self.assertEqual(status["status"], NER_STATUS_DISABLED)
        self.assertEqual(status["enabled"], False)

    def test_ner_detection_handles_missing_spacy_dependency(self) -> None:
        with patch("ner._spacy_module", return_value=None):
            status = detect_ner_support(enabled=True)

        self.assertEqual(status["status"], NER_STATUS_DEPENDENCY_MISSING)
        self.assertEqual(status["model_name"], DEFAULT_NER_MODEL)

    def test_ner_detection_handles_missing_polish_model(self) -> None:
        fake_spacy = FakeSpacy(error=OSError("model not found"))

        with patch("ner._spacy_module", return_value=fake_spacy):
            status = detect_ner_support(enabled=True)

        self.assertEqual(status["status"], NER_STATUS_MODEL_MISSING)
        self.assertEqual(status["warning"], "local NER model is missing")

    def test_entity_detection_maps_model_labels_to_internal_labels(self) -> None:
        model = FakeNerModel(
            [
                ("Person Example", "persName"),
                ("Example Org", "orgName"),
                ("Warsaw Example", "placeName"),
            ]
        )
        fake_spacy = FakeSpacy(model=model)

        with patch("ner._spacy_module", return_value=fake_spacy):
            context = prepare_ner_context(enabled=True)
            entities, counters = detect_entities(
                "Person Example works for Example Org in Warsaw Example.",
                context,
            )

        self.assertEqual([entity.label for entity in entities], [
            "NER_PERSON",
            "NER_ORG",
            "NER_LOCATION",
        ])
        self.assertEqual(counters["NER_PERSON"], 1)
        self.assertEqual(counters["NER_ORG"], 1)
        self.assertEqual(counters["NER_LOCATION"], 1)

    def test_ner_filters_public_version_and_scientific_false_positives(self) -> None:
        model = FakeNerModel(
            [
                ("Ministra Zdrowia", "persName"),
                ("Pa\u0144stwowa Inspekcja Sanitarna", "orgName"),
                ("Version 1.0, 2007", "persName"),
                ("Streptococcus pneumoniae", "persName"),
                ("Salmonella", "persName"),
                ("R\u00f3\u017ca", "persName"),
                ("Jan Kowalski", "persName"),
            ]
        )

        with patch("ner._spacy_module", return_value=FakeSpacy(model=model)):
            context = prepare_ner_context(enabled=True)
            entities, counters, exclusions, linebreak_count = detect_entities_with_details(
                (
                    "Ministra Zdrowia. Pa\u0144stwowa Inspekcja Sanitarna. "
                    "Version 1.0, 2007. Streptococcus pneumoniae. "
                    "Salmonella. R\u00f3\u017ca. Jan Kowalski."
                ),
                context,
            )

        self.assertEqual([entity.label for entity in entities], ["NER_PERSON"])
        self.assertEqual(counters["NER_PERSON"], 1)
        self.assertEqual(exclusions["PUBLIC_INSTITUTION_PHRASE"], 2)
        self.assertEqual(exclusions["VERSION_LIKE"], 1)
        self.assertEqual(exclusions["SCIENTIFIC_NAME"], 3)
        self.assertEqual(linebreak_count, 0)

    def test_ner_filters_polish_public_health_and_ordinary_false_positives(self) -> None:
        ordinary_terms = [
            "EPIDEMIOLOGICZNEJ",
            "\u017bYWNO\u015aCI\u0104",
            "PUBLICZNEJ",
            "wiotkie",
            "PESEL",
            "NIP",
            "REGON",
        ]
        public_terms = [
            "Pa\u0144stwow\u0105 Inspekcj\u0119 Sanitarn\u0105",
            "Powiatowej Stacji",
            "Stanu Sanitarnego Powiatu",
            "Urz\u0119dem Statystycznym w Krakowie",
            "Wojew\u00f3dzkiej Stacji Sanitarno-Epidemiologicznej",
            "GUS",
            "ECDC",
            "\u015awiatow\u0105 Organizacj\u0119 Zdrowia",
            "WHO",
        ]
        health_terms = [
            "Salmonella Enteritidis",
            "WZW B",
            "B\u0142onica-T\u0119\u017cec-Krztusiec",
            "POLIO",
            "Haemophilus",
            "B\u0141ONICA",
            "T\u0119\u017cec-B\u0142onica",
            "Zaka\u017cenia Neisseria",
        ]
        rules = [(term, "persName") for term in ordinary_terms]
        rules.extend((term, "orgName") for term in public_terms)
        rules.extend((term, "persName") for term in health_terms)
        rules.extend(
            [
                ("Jan Kowalski", "persName"),
                ("Anna Nowak", "persName"),
            ]
        )
        text = ". ".join(ordinary_terms + public_terms + health_terms)
        text = f"{text}. Jan Kowalski. Anna Nowak."

        with patch("ner._spacy_module", return_value=FakeSpacy(FakeNerModel(rules))):
            context = prepare_ner_context(enabled=True)
            entities, counters, exclusions, linebreak_count = detect_entities_with_details(
                text,
                context,
            )

        self.assertEqual([entity.label for entity in entities], [
            "NER_PERSON",
            "NER_PERSON",
        ])
        self.assertEqual(counters["NER_PERSON"], 2)
        self.assertEqual(exclusions["ORDINARY_WORD"], len(ordinary_terms))
        self.assertEqual(exclusions["PUBLIC_INSTITUTION_PHRASE"], len(public_terms))
        self.assertEqual(exclusions["SCIENTIFIC_NAME"], len(health_terms))
        self.assertEqual(linebreak_count, 0)

    def test_ner_allowlist_normalizes_case_dashes_boundaries_and_pdf_partial_term(self) -> None:
        public_terms = [
            "Wojew\u00f3dzkiej Stacji Sanitarno\u2013Epidemiologicznej",
            "Wojew\u00f3dzka Stacja Sanitarno \u2014 Epidemiologiczna",
            "ojew\u00f3dzkiej Stacji Sanitarno-Epidemiologicznej",
            "Raport: Wojew\u00f3dzkiej Stacji Sanitarno-Epidemiologicznej",
        ]
        health_terms = [
            "Odra",
            "WZW B",
            "POLIO",
            "Haemophilus influenzae",
            "Neisseria meningitidis",
            "B\u0141ONICA",
            "T\u0119\u017cec\u2013B\u0142onica",
            "B\u0142onica\u2014T\u0119\u017cec\u2013Krztusiec",
        ]
        model_public_terms = [
            "Wojew\u00f3dzkiej Stacji Sanitarno-Epidemiologicznej",
            "Wojew\u00f3dzka Stacja Sanitarno - Epidemiologiczna",
            "ojew\u00f3dzkiej Stacji Sanitarno-Epidemiologicznej",
            "Raport: Wojew\u00f3dzkiej Stacji Sanitarno-Epidemiologicznej",
        ]
        model_health_terms = [
            "Odra",
            "WZW B",
            "POLIO",
            "Haemophilus influenzae",
            "Neisseria meningitidis",
            "B\u0141ONICA",
            "T\u0119\u017cec-B\u0142onica",
            "B\u0142onica-T\u0119\u017cec-Krztusiec",
        ]
        rules = [(term, "orgName") for term in model_public_terms]
        rules.extend((term, "persName") for term in model_health_terms)
        rules.append(("Jan Kowalski", "persName"))
        text = ". ".join(public_terms + health_terms + ["Jan Kowalski"])

        with patch("ner._spacy_module", return_value=FakeSpacy(FakeNerModel(rules))):
            context = prepare_ner_context(enabled=True)
            entities, counters, exclusions, linebreak_count = detect_entities_with_details(
                text,
                context,
            )

        self.assertEqual([entity.label for entity in entities], ["NER_PERSON"])
        self.assertEqual(counters["NER_PERSON"], 1)
        self.assertEqual(exclusions["PUBLIC_INSTITUTION_PHRASE"], len(public_terms))
        self.assertEqual(exclusions["SCIENTIFIC_NAME"], len(health_terms))
        self.assertEqual(linebreak_count, 0)

    def test_txt_flow_keeps_false_positive_terms_but_redacts_strong_values(self) -> None:
        model = FakeNerModel(
            [
                ("EPIDEMIOLOGICZNEJ", "persName"),
                ("\u017bYWNO\u015aCI\u0104", "persName"),
                ("PUBLICZNEJ", "persName"),
                ("wiotkie", "persName"),
                ("WHO", "orgName"),
                ("GUS", "orgName"),
                ("WZW B", "orgName"),
                ("POLIO", "placeName"),
                ("Haemophilus", "orgName"),
                ("B\u0141ONICA", "placeName"),
                ("T\u0119\u017cec-B\u0142onica", "MISC"),
                ("Jan Kowalski", "persName"),
                ("Anna Nowak", "persName"),
                ("Jan Testowy", "persName"),
            ]
        )

        with workspace_temp_dir() as temp_dir:
            source_path = Path(temp_dir) / "document.txt"
            dictionary_path = Path(temp_dir) / "sensitive_terms.txt"
            dictionary_path.write_text("WHO = [PRIVATE_ORG]\n", encoding="utf-8")
            source_path.write_text(
                (
                    "EPIDEMIOLOGICZNEJ \u017bYWNO\u015aCI\u0104 PUBLICZNEJ wiotkie. "
                    "GUS WZW B POLIO Haemophilus B\u0141ONICA T\u0119\u017cec-B\u0142onica. WHO. "
                    "Jan Kowalski. Anna Nowak. mgr in\u017c. Jan Testowy. "
                    "Email safe@example.test. PESEL 00000000000. tel. 123 456 789."
                ),
                encoding="utf-8",
            )

            with patch("ner._spacy_module", return_value=FakeSpacy(model)):
                output_path, counters = anonymize_file(
                    source_path,
                    sensitive_terms_path=dictionary_path,
                    use_ner=True,
                )

            output_text = output_path.read_text(encoding="utf-8")
            report_text = (Path(temp_dir) / "document_RAPORT.txt").read_text(
                encoding="utf-8"
            )

        for safe_term in (
            "EPIDEMIOLOGICZNEJ",
            "\u017bYWNO\u015aCI\u0104",
            "PUBLICZNEJ",
            "wiotkie",
            "GUS",
            "WZW B",
            "POLIO",
            "Haemophilus",
            "B\u0141ONICA",
            "T\u0119\u017cec-B\u0142onica",
        ):
            self.assertIn(safe_term, output_text)
        self.assertIn("[PRIVATE_ORG]", output_text)
        self.assertEqual(output_text.count("[NER_PERSON]"), 3)
        self.assertIn("mgr in\u017c. [NER_PERSON]", output_text)
        self.assertIn("[EMAIL]", output_text)
        self.assertIn("[PESEL]", output_text)
        self.assertIn("[TELEFON]", output_text)
        self.assertEqual(counters["NER_PERSON"], 3)
        self.assertEqual(counters["PRIVATE_ORG"], 1)
        self.assertEqual(counters["EMAIL"], 1)
        self.assertEqual(counters["PESEL"], 1)
        self.assertEqual(counters["TELEFON"], 1)
        self.assertIn("* ORDINARY_WORD: 4", report_text)
        self.assertIn("* PUBLIC_INSTITUTION_PHRASE: 1", report_text)
        self.assertIn("* SCIENTIFIC_NAME: 5", report_text)
        self.assertNotIn("WHO", report_text)
        self.assertNotIn("Jan Kowalski", report_text)

    def test_ner_detects_person_across_soft_line_break(self) -> None:
        model = FakeNerModel([("Jan Kowalski", "persName")])

        with patch("ner._spacy_module", return_value=FakeSpacy(model=model)):
            context = prepare_ner_context(enabled=True)
            entities, counters, exclusions, linebreak_count = detect_entities_with_details(
                "Uczestnik Jan\nKowalski startowa\u0142 w konkursie.",
                context,
            )

        self.assertEqual(len(entities), 1)
        self.assertEqual(counters["NER_PERSON"], 1)
        self.assertEqual(sum(exclusions.values()), 0)
        self.assertEqual(linebreak_count, 1)

    def test_ner_does_not_join_unrelated_lowercase_line_break(self) -> None:
        model = FakeNerModel([("sanitarno epidemiologiczna", "persName")])

        with patch("ner._spacy_module", return_value=FakeSpacy(model=model)):
            context = prepare_ner_context(enabled=True)
            entities, counters, _, linebreak_count = detect_entities_with_details(
                "sanitarno\nepidemiologiczna stacja",
                context,
            )

        self.assertEqual(entities, [])
        self.assertEqual(counters["NER_PERSON"], 0)
        self.assertEqual(linebreak_count, 0)

    def test_ner_skips_non_person_entity_that_only_exists_via_line_break_bridging(
        self,
    ) -> None:
        """Regression test for a pilot finding: the soft line-break bridging
        built for split person names (see
        test_ner_detects_person_across_soft_line_break) does not know in
        advance what an entity will turn out to be, so it can also bridge
        two unrelated capitalized words from separate lines - for example
        two section headers - into what the model reports as a single
        organization. Such a non-person entity must be skipped rather than
        redacted, while a genuine split person name must still work."""
        model = FakeNerModel([("Rozdzial Pierwszy Podsumowanie", "orgName")])

        with patch("ner._spacy_module", return_value=FakeSpacy(model=model)):
            context = prepare_ner_context(enabled=True)
            entities, counters, exclusions, linebreak_count = (
                detect_entities_with_details(
                    "Rozdzial Pierwszy\nPodsumowanie wynikow badania.",
                    context,
                )
            )

        self.assertEqual(entities, [])
        self.assertEqual(counters["NER_ORG"], 0)
        self.assertEqual(exclusions[NER_EXCLUSION_LINEBREAK_NON_PERSON], 1)
        self.assertEqual(linebreak_count, 0)

    def test_person_left_expansion_masks_simple_capitalized_previous_token(self) -> None:
        person_first = "Jan"
        person_last = "Kowalski"
        model = FakeNerModel([(person_last, "persName")])

        with workspace_temp_dir() as temp_dir:
            source_path = Path(temp_dir) / "document.txt"
            source_path.write_text(
                f"{person_first} {person_last} wrote a note.",
                encoding="utf-8",
            )

            with patch("ner._spacy_module", return_value=FakeSpacy(model=model)):
                output_path, counters = anonymize_file(source_path, use_ner=True)

            output_text = output_path.read_text(encoding="utf-8")
            report_text = (Path(temp_dir) / "document_RAPORT.txt").read_text(
                encoding="utf-8"
            )

        self.assertEqual(output_text, "[NER_PERSON] wrote a note.")
        self.assertEqual(counters["NER_PERSON"], 1)
        self.assertNotIn(person_first, report_text)
        self.assertNotIn(person_last, report_text)

    def test_person_left_expansion_does_not_cross_lowercase_previous_word(self) -> None:
        model = FakeNerModel([("Kowalski", "persName")])

        with workspace_temp_dir() as temp_dir:
            source_path = Path(temp_dir) / "document.txt"
            source_path.write_text("pan Kowalski wrote a note.", encoding="utf-8")

            with patch("ner._spacy_module", return_value=FakeSpacy(model=model)):
                output_path, counters = anonymize_file(source_path, use_ner=True)

            output_text = output_path.read_text(encoding="utf-8")

        self.assertEqual(output_text, "pan [NER_PERSON] wrote a note.")
        self.assertEqual(counters["NER_PERSON"], 1)

    def test_single_token_person_with_title_context_is_kept(self) -> None:
        model = FakeNerModel([("Kowalski", "persName")])

        with workspace_temp_dir() as temp_dir:
            source_path = Path(temp_dir) / "document.txt"
            source_path.write_text("dr Kowalski wrote a note.", encoding="utf-8")

            with patch("ner._spacy_module", return_value=FakeSpacy(model=model)):
                output_path, counters = anonymize_file(source_path, use_ner=True)

            output_text = output_path.read_text(encoding="utf-8")

        self.assertEqual(output_text, "dr [NER_PERSON] wrote a note.")
        self.assertEqual(counters["NER_PERSON"], 1)

    def test_person_left_expansion_does_not_cross_placeholder(self) -> None:
        model = FakeNerModel([("Kowalski", "persName")])

        with workspace_temp_dir() as temp_dir:
            source_path = Path(temp_dir) / "document.txt"
            source_path.write_text("[EMAIL] Kowalski wrote a note.", encoding="utf-8")

            with patch("ner._spacy_module", return_value=FakeSpacy(model=model)):
                output_path, counters = anonymize_file(source_path, use_ner=True)

            output_text = output_path.read_text(encoding="utf-8")

        self.assertEqual(output_text, "[EMAIL] Kowalski wrote a note.")
        self.assertEqual(counters, {})

    def test_person_left_expansion_does_not_cross_punctuation(self) -> None:
        model = FakeNerModel([("Kowalski", "persName")])

        with workspace_temp_dir() as temp_dir:
            source_path = Path(temp_dir) / "document.txt"
            source_path.write_text("Jan, Kowalski wrote a note.", encoding="utf-8")

            with patch("ner._spacy_module", return_value=FakeSpacy(model=model)):
                output_path, counters = anonymize_file(source_path, use_ner=True)

            output_text = output_path.read_text(encoding="utf-8")

        self.assertEqual(output_text, "Jan, Kowalski wrote a note.")
        self.assertEqual(counters, {})

    def test_person_left_expansion_does_not_cross_all_caps_token(self) -> None:
        model = FakeNerModel([("Kowalski", "persName")])

        with workspace_temp_dir() as temp_dir:
            source_path = Path(temp_dir) / "document.txt"
            source_path.write_text("ACME Kowalski wrote a note.", encoding="utf-8")

            with patch("ner._spacy_module", return_value=FakeSpacy(model=model)):
                output_path, counters = anonymize_file(source_path, use_ner=True)

            output_text = output_path.read_text(encoding="utf-8")

        self.assertEqual(output_text, "ACME Kowalski wrote a note.")
        self.assertEqual(counters, {})

    def test_person_left_expansion_does_not_cross_line_break(self) -> None:
        model = FakeNerModel([("Kowalski", "persName")])

        with workspace_temp_dir() as temp_dir:
            source_path = Path(temp_dir) / "document.txt"
            source_path.write_text("Jan\nKowalski wrote a note.", encoding="utf-8")

            with patch("ner._spacy_module", return_value=FakeSpacy(model=model)):
                output_path, counters = anonymize_file(source_path, use_ner=True)

            output_text = output_path.read_text(encoding="utf-8")

        self.assertEqual(output_text, "Jan\nKowalski wrote a note.")
        self.assertEqual(counters, {})

    def test_existing_full_person_detection_still_masks_original_span(self) -> None:
        person = "Anna Nowak"
        model = FakeNerModel([(person, "persName")])

        with workspace_temp_dir() as temp_dir:
            source_path = Path(temp_dir) / "document.txt"
            source_path.write_text(f"{person} wrote a note.", encoding="utf-8")

            with patch("ner._spacy_module", return_value=FakeSpacy(model=model)):
                output_path, counters = anonymize_file(source_path, use_ner=True)

            output_text = output_path.read_text(encoding="utf-8")

        self.assertEqual(output_text, "[NER_PERSON] wrote a note.")
        self.assertEqual(counters["NER_PERSON"], 1)

    def test_regex_masks_email_phone_and_pesel_before_ner(self) -> None:
        model = FakeNerModel(
            [
                ("Kowalski", "persName"),
                ("[EMAIL]", "persName"),
                ("[TELEFON]", "persName"),
                ("[PESEL]", "persName"),
            ]
        )

        with workspace_temp_dir() as temp_dir:
            source_path = Path(temp_dir) / "document.txt"
            source_path.write_text(
                "Jan Kowalski, email safe@example.test, "
                "tel. 501 222 333, PESEL 90010112345.",
                encoding="utf-8",
            )

            with patch("ner._spacy_module", return_value=FakeSpacy(model=model)):
                output_path, counters = anonymize_file(source_path, use_ner=True)

            output_text = output_path.read_text(encoding="utf-8")
            report_text = (Path(temp_dir) / "document_RAPORT.txt").read_text(
                encoding="utf-8"
            )

        self.assertIn("[NER_PERSON]", output_text)
        self.assertIn("[EMAIL]", output_text)
        self.assertIn("[TELEFON]", output_text)
        self.assertIn("[PESEL]", output_text)
        self.assertEqual(counters["NER_PERSON"], 1)
        self.assertEqual(counters["EMAIL"], 1)
        self.assertEqual(counters["TELEFON"], 1)
        self.assertEqual(counters["PESEL"], 1)
        for source_value in ("safe@example.test", "501 222 333", "90010112345"):
            self.assertNotIn(source_value, output_text)
            self.assertNotIn(source_value, report_text)

    def test_txt_workflow_anonymizes_person_org_and_location_with_safe_report(self) -> None:
        person = "Person Example"
        org = "Example Org"
        location = "Warsaw Example"
        model = FakeNerModel(
            [
                (person, "PERSON"),
                (org, "ORG"),
                (location, "GPE"),
            ]
        )

        with workspace_temp_dir() as temp_dir:
            source_path = Path(temp_dir) / "document.txt"
            source_path.write_text(
                f"{person} met {org} in {location}. Contact safe@example.test.",
                encoding="utf-8",
            )

            with patch("ner._spacy_module", return_value=FakeSpacy(model=model)):
                output_path, counters = anonymize_file(source_path, use_ner=True)

            report_text = (Path(temp_dir) / "document_RAPORT.txt").read_text(
                encoding="utf-8"
            )
            output_text = output_path.read_text(encoding="utf-8")

        self.assertIn("[NER_PERSON]", output_text)
        self.assertIn("[NER_ORG]", output_text)
        self.assertIn("[NER_LOCATION]", output_text)
        self.assertIn("[EMAIL]", output_text)
        self.assertEqual(counters["NER_PERSON"], 1)
        self.assertEqual(counters["NER_ORG"], 1)
        self.assertEqual(counters["NER_LOCATION"], 1)
        self.assertEqual(counters["EMAIL"], 1)
        self.assertIn("NER enabled: yes", report_text)
        self.assertIn("NER status: available", report_text)
        self.assertIn("* NER_PERSON: 1", report_text)
        for source_value in (person, org, location, "safe@example.test"):
            self.assertNotIn(source_value, report_text)

    def test_ner_skips_existing_dictionary_and_regex_placeholders(self) -> None:
        private_person = "Person Private"
        org = "Example Org"
        model = FakeNerModel(
            [
                (private_person, "PERSON"),
                ("[EMAIL]", "ORG"),
                (org, "ORG"),
            ]
        )

        with workspace_temp_dir() as temp_dir:
            dictionary_path = Path(temp_dir) / "sensitive_terms.txt"
            dictionary_path.write_text(
                f"{private_person} = [PRIVATE_PERSON]\n",
                encoding="utf-8",
            )
            source_path = Path(temp_dir) / "document.txt"
            source_path.write_text(
                f"{private_person} contacted {org} at safe@example.test.",
                encoding="utf-8",
            )

            with patch("ner._spacy_module", return_value=FakeSpacy(model=model)):
                output_path, counters = anonymize_file(
                    source_path,
                    sensitive_terms_path=dictionary_path,
                    use_ner=True,
                )

            output_text = output_path.read_text(encoding="utf-8")

        self.assertIn("[PRIVATE_PERSON]", output_text)
        self.assertIn("[NER_ORG]", output_text)
        self.assertIn("[EMAIL]", output_text)
        self.assertNotIn("[NER_PERSON]", output_text)
        self.assertEqual(counters["PRIVATE_PERSON"], 1)
        self.assertEqual(counters["NER_ORG"], 1)
        self.assertEqual(counters["EMAIL"], 1)

    def test_missing_ner_dependency_does_not_crash_or_change_output(self) -> None:
        with workspace_temp_dir() as temp_dir:
            source_path = Path(temp_dir) / "document.txt"
            source_path.write_text(
                "Person Example contacted safe@example.test.",
                encoding="utf-8",
            )

            with patch("ner._spacy_module", return_value=None):
                output_path, _ = anonymize_file(
                    source_path,
                    use_ner=True,
                )

            output_text = output_path.read_text(encoding="utf-8")
            report_text = (Path(temp_dir) / "document_RAPORT.txt").read_text(
                encoding="utf-8"
            )

        self.assertIn("Person Example", output_text)
        self.assertIn("[EMAIL]", output_text)
        self.assertIn("NER status: dependency_missing", report_text)
        self.assertNotIn("safe@example.test", report_text)

    def test_batch_summary_contains_safe_ner_status_and_counters(self) -> None:
        person = "Person Example"
        model = FakeNerModel([(person, "PERSON")])

        with workspace_temp_dir() as temp_dir:
            source_dir = Path(temp_dir) / "source"
            output_dir = Path(temp_dir) / "output"
            source_dir.mkdir()
            output_dir.mkdir()
            source_path = source_dir / "document.txt"
            source_path.write_text(f"{person} wrote a note.", encoding="utf-8")

            with patch("ner._spacy_module", return_value=FakeSpacy(model=model)):
                result = anonymize_batch([source_path], output_dir, use_ner=True)

            summary_text = result.summary_path.read_text(encoding="utf-8")

        self.assertEqual(result.ner_status_counts["available"], 1)
        self.assertEqual(result.ner_category_counters["NER_PERSON"], 1)
        self.assertIn("Local NER:", summary_text)
        self.assertIn("* files processed with NER: 1", summary_text)
        self.assertIn("* NER_PERSON: 1", summary_text)
        self.assertIn("NER status: available", summary_text)
        self.assertNotIn(person, summary_text)
        self.assertNotIn(str(source_dir), summary_text)

    def test_docx_workflow_can_use_ner_without_requiring_real_model(self) -> None:
        person = "Person Example"
        model = FakeNerModel([(person, "PERSON")])

        with workspace_temp_dir() as temp_dir:
            source_path = Path(temp_dir) / "document.docx"
            document = Document()
            document.add_paragraph(f"{person} wrote a note.")
            document.save(source_path)

            with patch("ner._spacy_module", return_value=FakeSpacy(model=model)):
                output_path, counters = anonymize_docx_file(
                    source_path,
                    use_ner=True,
                )

            report_text = (Path(temp_dir) / "document_RAPORT.txt").read_text(
                encoding="utf-8"
            )

        self.assertEqual(output_path.name, "document_ANON.docx")
        self.assertEqual(counters["NER_PERSON"], 1)
        self.assertIn("* NER_PERSON: 1", report_text)
        self.assertNotIn(person, report_text)


if __name__ == "__main__":
    unittest.main()
